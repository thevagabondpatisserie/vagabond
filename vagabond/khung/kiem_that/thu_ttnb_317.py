"""#317: lô TTNB trên DB thật, hai người nhận là hai giao dịch khác nhau."""
import base64
import io
from unittest.mock import patch
import frappe
from openpyxl import load_workbook
from vagabond import de_nghi_chi as dc, ttnb_lo
from vagabond.khung.kiem_that.nen import ca, la, dung, _DA_TAO
from vagabond.khung.kiem_that.thu_ho_so_tt_v445 import _giao_dich_ngan_hang


@ca('#317 lô TTNB: ba phiếu hai tài khoản, sao kê thật, retry và Excel')
def _lo_that():
	ct = frappe.db.get_single_value('Global Defaults', 'default_company')
	ds=[]
	for i in range(3):
		d=frappe.new_doc(dc.DT)
		d.ten_khoan_chi='KIEM317-' + frappe.generate_hash(length=8)
		d.company=ct
		d.loai_nghiep_vu='Tạm ứng'
		d.nguoi_tao=frappe.session.user
		d.phuong_thuc='Chuyển khoản'
		d.hinh_thuc='Hoàn tiền cho nhân viên'
		d.ten_tk='NGUOI THU %s' % (1 if i<2 else 2)
		d.so_tk='3170001' if i<2 else '3170002'
		d.ngan_hang='MB'
		d.append('cac_khoan',dict(noi_dung='Kiểm lô',so_tien=100000))
		d.insert(ignore_permissions=True)
		_DA_TAO.append((d.doctype,d.name))
		# Fixture bắt đầu sau duyệt, không nhận là kiểm toàn chuỗi duyệt.
		frappe.db.set_value(dc.DT,d.name,'trang_thai',dc.TT_CHO_KE_TOAN)
		ds.append(d.name)
	lo=ttnb_lo.gop(ds,1)['lo']
	la('hai người nhận hai lô',len(lo),2)
	la('retry giữ lô',ttnb_lo.gop(list(reversed(ds)),1)['lo'],lo)
	g=lo[0]
	bt=_giao_dich_ngan_hang(g['ma_lo'],g['tong_tien'],ct,noi_dung=g['noi_dung'])
	ttnb_lo.khop(bt.as_dict())
	for ma in ds:
		d=frappe.get_doc(dc.DT,ma)
		la('chỉ đúng lô được tất toán',d.trang_thai,dc.TT_DA_CHI if ma in g['phieu'] else dc.TT_CHO_KE_TOAN)
	ttnb_lo.khop(bt.as_dict())
	for g in lo[1:]:
		bt=_giao_dich_ngan_hang(g['ma_lo'],g['tong_tien'],ct,noi_dung=g['noi_dung'])
		ttnb_lo.khop(bt.as_dict())
	la('cả ba chi đủ',frappe.db.count(dc.DT,{'name':['in',ds],'trang_thai':dc.TT_DA_CHI}),3)
	r=dc.xuat_excel(tim='KIEM317-',so_ngay=1)
	wb=load_workbook(io.BytesIO(base64.b64decode(r['b64'])))
	la('đủ cột',wb.active.max_column,19)
	dung('có ba phiếu trong Excel',set(ds).issubset({r[0] for r in wb.active.iter_rows(min_row=2,values_only=True)}))
